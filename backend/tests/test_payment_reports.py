import asyncio
import csv
import json
import os
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from io import BytesIO, StringIO
from unittest.mock import patch
from urllib.parse import urlsplit
from threading import Barrier
from uuid import uuid4
from xml.etree.ElementTree import fromstring

from fastapi import FastAPI, Header, HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from database import Base
from models import AutopayReceipt, PremiumOrder, User
from payment_reports import (
    build_payments_csv, build_payments_router, build_payments_workbook,
    monthly_totals, payment_overview, receipt_query, record_autopay_receipt, year_query,
)
from premium_payments import payment_hash, process_itn


class PaymentReportsTests(unittest.TestCase):
    def make_engine(self):
        return create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)

    def setUp(self):
        self.engine = self.make_engine()
        if self.engine.dialect.name == "sqlite":
            with self.engine.connect() as connection:
                connection.exec_driver_sql("PRAGMA foreign_keys=ON")
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine, autoflush=False)()
        self.user = User(email="buyer@example.test", first_name="Łucja", last_name="Żółć", is_active=1)
        self.db.add(self.user)
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def order(self, amount="10.00", environment="production", package="organizer", period="yearly"):
        order = PremiumOrder(
            id=uuid4().hex, user_id=self.user.id, package=package, period=period,
            amount=amount, currency="PLN", environment=environment, service_id="12345",
            status="pending", created_at="2026-10-25T12:00:00Z", terms_version="test",
            terms_accepted_at="2026-10-25T12:00:00Z", offer_snapshot="{}",
        )
        self.db.add(order)
        self.db.commit()
        return order

    def transaction(self, order, remote="remote1", paid="20261025140000", status="SUCCESS"):
        return fromstring(f"<transaction><orderID>{order.id}</orderID><remoteID>{remote}</remoteID>"
                          f"<amount>{order.amount}</amount><currency>PLN</currency>"
                          f"<paymentDate>{paid}</paymentDate><paymentStatus>{status}</paymentStatus></transaction>")

    def receipt(self, amount="10.00", remote="remote1", paid="20261025140000", environment="production"):
        order = self.order(amount=amount, environment=environment)
        receipt = record_autopay_receipt(order, self.transaction(order, remote, paid), self.db)
        self.db.commit()
        return receipt

    def itn(self, order, remote="remote1", status="SUCCESS", paid="20261025140000", tamper=False):
        config = {"service": "12345", "key": "test-key", "algorithm": "sha256"}
        transaction = self.transaction(order, remote, paid, status)
        signature = payment_hash(["12345"] + [element.text for element in transaction], config)
        from xml.etree.ElementTree import tostring
        xml = (b"<transactionList><serviceID>12345</serviceID><transactions>" + tostring(transaction)
               + b"</transactions><hash>" + signature.encode() + b"</hash></transactionList>")
        if tamper:
            xml = xml.replace(b"SUCCESS", b"FAILURE")
        return process_itn(xml, self.db, config)

    def test_exact_sums_and_environment_separation(self):
        self.receipt("0.10", "one")
        self.receipt("0.20", "two")
        self.receipt("999.99", "sandbox", environment="sandbox")
        overview = payment_overview(self.db, 2026, "production")
        self.assertEqual(overview["annual"], {"count": 2, "amount": "0.30"})
        self.assertEqual(overview["lifetime"], overview["annual"])
        self.assertEqual(overview["months"][9]["amount"], "0.30")
        self.assertEqual(payment_overview(self.db, 2026, "sandbox")["annual"]["amount"], "999.99")

    def test_payment_year_uses_provider_date_in_poland_not_notification_arrival(self):
        december = self.receipt("10.00", "dec", paid="20251231235959")
        january = self.receipt("20.00", "jan", paid="20260101000001")
        self.assertEqual(december.paid_on, date(2025, 12, 31))
        self.assertEqual(january.paid_on, date(2026, 1, 1))
        self.assertIn("+01:00", january.paid_at)
        self.assertEqual(payment_overview(self.db, 2025, "production")["annual"]["amount"], "10.00")
        self.assertEqual(payment_overview(self.db, 2026, "production")["annual"]["amount"], "20.00")

    def test_duplicate_notification_counts_once_but_two_payments_for_one_order_count_twice(self):
        order = self.order(environment="sandbox")
        for _ in range(3):
            self.itn(order)
        self.assertEqual(self.db.query(AutopayReceipt).count(), 1)
        self.itn(order, remote="remote2")
        self.assertEqual(self.db.query(AutopayReceipt).count(), 2)
        self.assertEqual(payment_overview(self.db, 2026, "sandbox")["annual"]["amount"], "20.00")
        self.assertEqual(order.status, "paid")
        self.assertIsNone(self.user.premium_organizer_paid_until)

    def test_pending_failure_bad_hash_and_unknown_order_do_not_create_receipts(self):
        order = self.order(environment="sandbox")
        self.itn(order, status="PENDING")
        self.itn(order, status="FAILURE")
        with self.assertRaises(HTTPException):
            self.itn(order, tamper=True)
        unknown = self.order(environment="sandbox")
        self.db.delete(unknown)
        self.db.commit()
        reply = self.itn(unknown, remote="unknown")
        self.assertEqual(fromstring(reply).findtext(".//confirmation"), "NOTCONFIRMED")
        self.assertEqual(self.db.query(AutopayReceipt).count(), 0)

    def test_inconsistent_reuse_of_autopay_transaction_id_is_rejected(self):
        self.receipt(remote="unique")
        other = self.order(amount="20.00")
        with self.assertRaises(HTTPException) as error:
            record_autopay_receipt(other, self.transaction(other, remote="unique"), self.db)
        self.assertEqual(error.exception.status_code, 409)
        self.assertEqual(self.db.query(AutopayReceipt).count(), 1)

    def test_invalid_date_or_fractional_cent_does_not_create_receipt(self):
        order = self.order()
        for paid in ("20260230000000", "bad", "19991231000000"):
            with self.subTest(paid=paid), self.assertRaises(HTTPException):
                record_autopay_receipt(order, self.transaction(order, paid=paid), self.db)
        transaction = self.transaction(order)
        transaction.find("amount").text = "0.001"
        with self.assertRaises(HTTPException):
            record_autopay_receipt(order, transaction, self.db)
        self.assertEqual(self.db.query(AutopayReceipt).count(), 0)

    def test_filters_and_pagination_do_not_change_annual_summary(self):
        self.receipt("10.00", "january", paid="20260110120000")
        self.receipt("20.00", "october", paid="20261010120000")
        result = payment_overview(self.db, 2026, "production", month=10, page_size=1)
        self.assertEqual(result["matching"], {"count": 1, "amount": "20.00"})
        self.assertEqual(result["annual"]["amount"], "30.00")
        self.assertEqual(result["entries"][0]["remote_id"], "october")
        self.assertEqual(len(payment_overview(self.db, 2026, "production", page=2, page_size=1)["entries"]), 1)
        self.assertEqual(payment_overview(self.db, 2026, "production", search="buyer@")["matching"]["count"], 2)
        self.assertEqual(payment_overview(self.db, 2026, "production", search="%")["matching"]["count"], 0)

    def test_receipt_keeps_payer_and_order_reference_after_account_and_order_deletion(self):
        entry = self.receipt()
        order = self.db.get(PremiumOrder, entry.order_id)
        reference = order.id
        self.db.delete(self.user)
        self.db.commit()
        self.db.delete(order)
        self.db.commit()
        self.db.refresh(entry)
        self.assertIsNone(entry.order_id)
        self.assertEqual(entry.order_reference, reference)
        self.assertEqual(entry.payer_email, "buyer@example.test")
        self.assertEqual(entry.payer_name, "Łucja Żółć")

    def test_excel_and_csv_include_full_year_and_never_evaluate_payer_formulas(self):
        self.user.first_name = '=HYPERLINK("https://example.invalid")'
        self.db.commit()
        self.receipt("10.10", "one")
        self.receipt("20.20", "two")
        self.receipt("999.99", "test", environment="sandbox")
        query = year_query(receipt_query(self.db, "production"), 2026)
        entries = query.order_by(AutopayReceipt.id).all()
        workbook = load_workbook(BytesIO(build_payments_workbook(entries, 2026, monthly_totals(query), {"name": "Piotr Bronny"})))
        self.assertEqual(workbook.sheetnames, ["Podsumowanie", "Miesiące", "Wpłaty"])
        self.assertEqual(workbook["Podsumowanie"]["B7"].value, 30.3)
        self.assertEqual(workbook["Wpłaty"].max_row, 3)
        self.assertEqual(workbook["Wpłaty"]["B2"].data_type, "s")
        self.assertTrue(workbook["Wpłaty"]["B2"].value.startswith("=HYPERLINK"))
        self.assertEqual(workbook["Miesiące"].max_row, 13)
        raw = build_payments_csv(entries, 2026)
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(StringIO(raw.decode("utf-8-sig")), delimiter=";"))
        self.assertTrue(rows[1][1].startswith("'=HYPERLINK"))
        self.assertEqual(rows[-2][4], "30,30")
        self.assertNotIn("999,99", raw.decode("utf-8-sig"))

    def test_empty_year_has_zero_summary_and_downloadable_workbook(self):
        result = payment_overview(self.db, 2026, "production")
        self.assertEqual(result["annual"], {"count": 0, "amount": "0.00"})
        workbook = load_workbook(BytesIO(build_payments_workbook([], 2026, result["months"], {})))
        self.assertEqual(workbook["Podsumowanie"]["B7"].value, 0)
        self.assertEqual(workbook["Wpłaty"].max_row, 1)

    def request(self, path, token=None):
        app = FastAPI()

        def admin(authorization: str | None = Header(default=None)):
            if authorization is None:
                raise HTTPException(401, "Wymagane logowanie")
            if authorization != "Bearer admin":
                raise HTTPException(403, "Brak uprawnień administratora")
            return self.user

        app.include_router(build_payments_router(lambda: self.db, admin, lambda: {"name": "Piotr Bronny"}))
        result = []
        parsed = urlsplit(path)

        async def run():
            async def receive():
                return {"type": "http.request", "body": b"", "more_body": False}

            async def send(message):
                result.append(message)

            headers = [(b"authorization", f"Bearer {token}".encode())] if token else []
            await app({"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": "GET",
                       "scheme": "https", "path": parsed.path, "raw_path": parsed.path.encode(),
                       "root_path": "", "query_string": parsed.query.encode(), "headers": headers,
                       "server": ("example.test", 443), "client": ("127.0.0.1", 1234)}, receive, send)

        asyncio.run(run())
        start = next(item for item in result if item["type"] == "http.response.start")
        body = b"".join(item.get("body", b"") for item in result if item["type"] == "http.response.body")
        return start["status"], dict(start["headers"]), body

    def test_http_routes_require_admin_and_export_excludes_sandbox(self):
        self.receipt("12.34", "real")
        self.receipt("99.99", "test", environment="sandbox")
        for path in ("/admin/payments?year=2026", "/admin/payments/report?year=2026&format=csv"):
            self.assertEqual(self.request(path)[0], 401)
            self.assertEqual(self.request(path, "organizer")[0], 403)
        code, headers, body = self.request("/admin/payments?year=2026", "admin")
        self.assertEqual(code, 200)
        self.assertEqual(headers[b"cache-control"], b"no-store")
        self.assertEqual(json.loads(body)["annual"]["amount"], "12.34")
        with patch("payment_reports.monthly_totals", side_effect=AssertionError("Export must not re-query a different monthly snapshot")):
            code, headers, body = self.request("/admin/payments/report?year=2026&format=xlsx", "admin")
        self.assertEqual(code, 200)
        self.assertIn(b"autopay-wplaty-2026.xlsx", headers[b"content-disposition"])
        self.assertEqual(load_workbook(BytesIO(body))["Podsumowanie"]["B7"].value, 12.34)
        self.assertEqual(self.request("/admin/payments/report?year=2026&environment=sandbox", "admin")[0], 422)

    def test_http_rejects_invalid_filters(self):
        for query in ("year=0", "month=13", "environment=all", "page=0", "page_size=10000"):
            self.assertEqual(self.request(f"/admin/payments?{query}", "admin")[0], 422)


@unittest.skipUnless(os.environ.get("TEST_POSTGRES_URL"), "TEST_POSTGRES_URL is not configured")
class PostgresPaymentReportsTests(PaymentReportsTests):
    def make_engine(self):
        self.schema = "receipts_test_" + uuid4().hex
        self.admin_engine = create_engine(os.environ["TEST_POSTGRES_URL"])
        with self.admin_engine.begin() as connection:
            connection.execute(text(f'CREATE SCHEMA "{self.schema}"'))
        return create_engine(os.environ["TEST_POSTGRES_URL"], connect_args={
            "options": f"-csearch_path={self.schema} -cstatement_timeout=10000",
        })

    def tearDown(self):
        super().tearDown()
        with self.admin_engine.begin() as connection:
            connection.execute(text(f'DROP SCHEMA "{self.schema}" CASCADE'))
        self.admin_engine.dispose()

    def concurrent_itns(self, remote_ids):
        order = self.order(environment="sandbox")
        order_id = order.id
        barrier = Barrier(len(remote_ids))
        sessions = sessionmaker(bind=self.engine, autoflush=False)
        config = {"service": "12345", "key": "test-key", "algorithm": "sha256"}

        def notify(remote_id):
            signature = payment_hash(["12345", order_id, remote_id, "10.00", "PLN", "20261025140000", "SUCCESS"], config)
            xml = (f"<transactionList><serviceID>12345</serviceID><transactions><transaction>"
                   f"<orderID>{order_id}</orderID><remoteID>{remote_id}</remoteID><amount>10.00</amount>"
                   f"<currency>PLN</currency><paymentDate>20261025140000</paymentDate>"
                   f"<paymentStatus>SUCCESS</paymentStatus></transaction></transactions><hash>{signature}</hash></transactionList>")
            with sessions() as db:
                cached = db.get(PremiumOrder, order_id)
                barrier.wait(timeout=10)
                reply = process_itn(xml.encode(), db, config)
                self.assertIsNotNone(cached)
                return fromstring(reply).findtext(".//confirmation")

        with ThreadPoolExecutor(max_workers=len(remote_ids)) as pool:
            replies = list(pool.map(notify, remote_ids))
        self.assertEqual(replies, ["CONFIRMED"] * len(remote_ids))
        self.db.expire_all()

    def test_concurrent_repeated_itn_is_counted_once(self):
        self.concurrent_itns(["same"] * 4)
        self.assertEqual(self.db.query(AutopayReceipt).count(), 1)
        self.assertEqual(payment_overview(self.db, 2026, "sandbox")["annual"]["amount"], "10.00")

    def test_concurrent_different_transactions_for_one_order_are_all_counted(self):
        self.concurrent_itns(["parallel1", "parallel2", "parallel3", "parallel4"])
        self.assertEqual(self.db.query(AutopayReceipt).count(), 4)
        self.assertEqual(payment_overview(self.db, 2026, "sandbox")["annual"]["amount"], "40.00")


if __name__ == "__main__":
    unittest.main()
