from io import BytesIO
import unittest
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from database import Base
from models import Competition, CompetitionParticipant, Discipline, DisciplineResult, ParticipantDiscipline, User
from results_excel import build_results_workbook


# main runs legacy backfills on import; isolate them from the configured database.
bootstrap_engine = create_engine("sqlite://")
Base.metadata.create_all(bootstrap_engine)
with patch("database.SessionLocal", sessionmaker(bind=bootstrap_engine)):
    import main
bootstrap_engine.dispose()


class ResultsWorkbookTests(unittest.TestCase):
    def workbook(self, categories):
        return load_workbook(BytesIO(build_results_workbook(categories)))

    def test_plain_tables_preserve_text_and_numeric_values(self):
        book = self.workbook([("Pistolet", {"shooters": [{
            "place": 1, "last_name": "Żółć", "first_name": "Łucja",
            "club": '=HYPERLINK("https://example.test")',
            "license_number": "00123/2026", "points": "93.50",
        }]})])
        sheet = book.active
        self.assertEqual(list(sheet.values), [
            ("Miejsce", "Nazwisko", "Imię", "Klub", "Numer licencji", "Wynik"),
            (1, "Żółć", "Łucja", '=HYPERLINK("https://example.test")', "00123/2026", 93.5),
        ])
        self.assertEqual(sheet["D2"].data_type, "s")
        self.assertEqual(sheet["F2"].data_type, "n")
        self.assertEqual(sheet.auto_filter.ref, "A1:F2")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertFalse(sheet.merged_cells)

    def test_long_duplicate_invalid_titles_and_empty_disciplines(self):
        names = ["Pistolet:/[]*?" + "a" * 40] * 3 + ["'''", "pistolet:/[]*?" + "A" * 40]
        book = self.workbook((name, {"shooters": []}) for name in names)
        self.assertEqual(len(book.worksheets), 5)
        self.assertEqual(len({name.casefold() for name in book.sheetnames}), 5)
        for sheet in book:
            self.assertLessEqual(len(sheet.title), 31)
            self.assertNotRegex(sheet.title, r"[\\/*?:\[\]]")
            self.assertEqual(sheet.max_row, 1)

    def test_dynamic_and_skeet_details_are_separate_numeric_columns(self):
        book = self.workbook([
            ("Dynamiczna", {"shooters": [{
                "dynamic_stage": True, "points": "75.25", "score_points": "75.25",
                "hit_factor": "4.4321", "time_seconds": "12.50", "hits": 15, "targets_count": 16,
            }]}),
            ("Skeet", {"shooters": [{"points": "85", "round_scores": [45, 40]}]}),
            ("Strzelba", {"shooters": [{
                "practical_shotgun": True, "points": "0", "score_points": "0",
                "hit_factor": "DQ", "disqualified": True, "disqualification_reason": "Przekroczony czas",
            }]}),
        ])
        dynamic = dict(zip(*list(book.worksheets[0].values)))
        self.assertEqual(dynamic["Wynik"], 75.25)
        self.assertEqual(dynamic["Hit Factor"], 4.4321)
        self.assertEqual(dynamic["Czas [s]"], 12.5)
        skeet = dict(zip(*list(book.worksheets[1].values)))
        self.assertEqual((skeet["Seria 1"], skeet["Seria 2"]), (45, 40))
        practical = dict(zip(*list(book.worksheets[2].values)))
        self.assertEqual(practical["Wynik"], 0)
        self.assertEqual(practical["Hit Factor"], "DQ")
        self.assertIs(practical["Dyskwalifikacja"], True)

    def test_competition_without_disciplines_is_valid_workbook(self):
        book = self.workbook([])
        self.assertEqual(book.sheetnames, ["Wyniki"])
        self.assertEqual(book.active.max_row, 1)


class ResultsExcelEndpointTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()
        self.organizer = User(email="organizer@example.test", role="organizer", roles="organizer")
        self.competition = Competition(
            id=1, name="Łódzkie zawody", date="2099-01-01", location="Łódź",
            status="completed", created_by=self.organizer.email,
        )
        self.db.add(self.competition)
        self.db.flush()
        self.db.add_all([
            Discipline(id=1, competition_id=1, name="Karabin", scoring_type="points", shots_count=10, display_order=1),
            Discipline(id=2, competition_id=1, name="Pistolet", scoring_type="points", shots_count=10, display_order=0),
            CompetitionParticipant(
                id=1, competition_id=1, user_email="lucja@example.test", first_name="Łucja", last_name="Żółć",
                license_number="00123", club="Klub", entry_type="shooter", checked_in=1, paid=1,
            ),
            CompetitionParticipant(
                id=2, competition_id=1, user_email="jan@example.test", first_name="Jan", last_name="Nowak",
                entry_type="shooter", checked_in=1, paid=1,
            ),
        ])
        self.db.flush()
        for participant_id in (1, 2):
            self.db.add(ParticipantDiscipline(participant_id=participant_id, discipline_id=2, ammo_type="own"))
            self.db.add(DisciplineResult(
                competition_id=1, discipline_id=2, participant_id=participant_id,
                judge_email="judge@example.test", points="93.50",
            ))
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def download(self, user=None):
        with patch("main.auto_complete_started_competitions"):
            return main.download_organizer_competition_results_excel(1, user=user or self.organizer, db=self.db)

    def test_exports_each_discipline_in_order_with_current_ranking(self):
        response = self.download()
        self.assertEqual(response.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertTrue(response.headers["content-disposition"].endswith('.xlsx"'))
        book = load_workbook(BytesIO(response.body))
        self.assertEqual(book.sheetnames, ["Pistolet", "Karabin"])
        rows = list(book["Pistolet"].values)
        self.assertEqual([row[0] for row in rows[1:]], [1, 1])
        self.assertEqual({row[1] for row in rows[1:]}, {"Nowak", "Żółć"})
        self.assertEqual([row[5] for row in rows[1:]], [93.5, 93.5])
        self.assertIn("00123", [row[4] for row in rows[1:]])
        self.assertEqual(book["Karabin"].max_row, 1)

    def test_export_rejects_foreign_organizer_but_allows_admin(self):
        other = User(email="other@example.test", role="organizer", roles="organizer")
        with self.assertRaises(HTTPException) as error:
            self.download(other)
        self.assertEqual(error.exception.status_code, 403)
        admin = User(email="admin@example.test", role="admin", roles="admin")
        self.assertEqual(self.download(admin).status_code, 200)

    def test_export_requires_completed_competition(self):
        for status in ("started", "draft", "published", "cancelled"):
            with self.subTest(status=status):
                self.competition.status = status
                self.db.commit()
                with self.assertRaises(HTTPException) as error:
                    self.download()
                self.assertEqual(error.exception.status_code, 400)
