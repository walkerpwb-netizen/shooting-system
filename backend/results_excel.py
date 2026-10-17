"""Plain, editable result tables for organizer XLSX downloads."""

from decimal import Decimal, InvalidOperation
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter


BASE_COLUMNS = [
    ("place", "Miejsce", True),
    ("last_name", "Nazwisko", False),
    ("first_name", "Imię", False),
    ("club", "Klub", False),
    ("license_number", "Numer licencji", False),
    ("points", "Wynik", True),
]
DYNAMIC_COLUMNS = [
    ("hit_factor", "Hit Factor", True),
    ("time_seconds", "Czas [s]", True),
    ("hits", "Trafienia", True),
    ("targets_count", "Liczba celów", True),
]
PRACTICAL_COLUMNS = [
    ("time_limit_seconds", "Limit czasu [s]", True),
    ("disqualified", "Dyskwalifikacja", False),
    ("disqualification_reason", "Powód dyskwalifikacji", False),
]


def numeric_value(value):
    if value is None or value == "":
        return None
    try:
        number = Decimal(str(value).replace(",", "."))
    except InvalidOperation:
        return value
    return number if number.is_finite() else str(value)


def sheet_title(name, used_titles):
    title = re.sub(r"[\\/*?:\[\]]", "-", ILLEGAL_CHARACTERS_RE.sub("", name or ""))
    title = title.strip().strip("'") or "Konkurencja"
    candidate = title[:31]
    suffix = 2
    while candidate.casefold() in used_titles:
        ending = f" ({suffix})"
        candidate = title[:31 - len(ending)] + ending
        suffix += 1
    used_titles.add(candidate.casefold())
    return candidate


def build_results_workbook(categories) -> bytes:
    """Consume (discipline name, result payload) pairs in competition order."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles = set()

    for name, payload in categories:
        shooters = payload.get("shooters", [])
        sheet = workbook.create_sheet(sheet_title(name, used_titles))
        columns = list(BASE_COLUMNS)
        if any(row.get("dynamic_stage") or row.get("practical_shotgun") for row in shooters):
            columns.extend(DYNAMIC_COLUMNS)
        if any(row.get("practical_shotgun") for row in shooters):
            columns.extend(PRACTICAL_COLUMNS)
        rounds_count = max((len(row.get("round_scores", [])) for row in shooters), default=0)
        columns.extend((f"round_{index}", f"Seria {index + 1}", True) for index in range(rounds_count))
        sheet.append([label for _, label, _ in columns])

        for row_index, shooter in enumerate(shooters, start=2):
            values = dict(shooter)
            values["points"] = shooter.get("score_points", shooter.get("points", ""))
            values.update({f"round_{index}": score for index, score in enumerate(shooter.get("round_scores", []))})
            for column_index, (key, _, numeric) in enumerate(columns, start=1):
                value = values.get(key, "")
                if numeric:
                    value = numeric_value(value)
                cell = sheet.cell(row=row_index, column=column_index)
                if isinstance(value, str):
                    # Names and licenses are literal text, including leading '=' or zeros.
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
                    cell.data_type = "s"
                else:
                    cell.value = value

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, (_, label, _) in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(16, len(label) + 2)

    if not workbook.worksheets:
        workbook.create_sheet("Wyniki").append([label for _, label, _ in BASE_COLUMNS])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
