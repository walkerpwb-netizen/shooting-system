"""Private Autopay receipt ledger and cash-receipt reports. No manual entries."""

import csv
import re
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from typing import Literal
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import extract, func, or_

from models import AutopayReceipt, User
from autopay_protocol import parse_payment_date
from payment_limits import (
    LIMIT_SOURCE, QUARTER_LIMITS_GROSZ, acquire_payment_limit_lock, payment_limit_state,
)

WARSAW = ZoneInfo("Europe/Warsaw")
TAX_SOURCE = "https://podatki.gov.pl/podatki-osobiste/pit/informacje-podstawowe/co-jest-opodatkowane/dochody-z-dzialalnosci-nierejestrowanej"
FEE_TARIFF = "Standard: 1,19% + 0,34 zł"
FEE_NOTE = (
    "Prowizje i kwoty po prowizji są szacunkowe: Standard 1,19% + 0,34 zł od każdej wpłaty, "
    "zaokrąglenie do grosza osobno dla każdej transakcji. Nie są potwierdzeniem pobrania opłaty ani wypłaty na konto. "
    "Koszty do PIT-36 ustal na podstawie rzeczywistych dokumentów Autopay."
)
LIMIT_NOTE = (
    "Limit dotyczy przychodu należnego przed prowizjami i innymi kosztami, również kwot jeszcze niezapłaconych. "
    "Ten licznik obejmuje tylko zapisane wpłaty Autopay według daty płatności. "
    "Nie obejmuje nieopłaconych należności, sprzedaży poza systemem ani korekt i zwrotów. "
    "Nie zastępuje pełnej ewidencji przychodu należnego."
)
REPORT_NOTE = (
    "Zestawienie pomocnicze do PIT-36: tylko potwierdzone wpłaty klientów przez Autopay w tym systemie. "
    "Wpłaty wykazano w pełnej kwocie, a szacunek prowizji osobno. Inne koszty i zwroty nie są uwzględnione. "
    "Uzgodnij z dokumentami operatora. Raport nie jest deklaracją ani wyliczeniem podatku."
)
MONTHS = ["Styczeń", "Luty", "Marzec", "Kwiecień", "Maj", "Czerwiec", "Lipiec", "Sierpień", "Wrzesień", "Październik", "Listopad", "Grudzień"]
HEADERS = ["Data wpłaty (Polska)", "Płatnik", "E-mail konta", "Usługa", "Kwota klienta [PLN]", "ID transakcji Autopay", "ID serwisu", "Numer zamówienia", "Data rejestracji w systemie (UTC)"]
HEADERS += ["Prowizja szacunkowa [PLN]", "Po prowizji — szacunek [PLN]", "Taryfa szacunku"]


def money(grosz):
    return f"{Decimal(int(grosz or 0)) / 100:.2f}"


def estimated_fee(grosz):
    # Integer arithmetic, ROUND_HALF_UP to a grosz for each transaction, never on a total.
    return (grosz * 119 + 5000) // 10000 + 34


def financial_totals(count, amount, fee):
    return {"count": count, "amount": money(amount), "estimated_fee": money(fee),
            "estimated_net": money(int(amount or 0) - int(fee or 0))}


def receipt_description(package, period):
    if package == "organizer":
        return "Roczny Organizator + Premium Strzelca" if period == "yearly" else "Jednorazowa dodatkowa publikacja"
    return "Premium Strzelca — rok" if period == "yearly" else "Premium Strzelca — miesiąc"


def record_autopay_receipt(order, transaction, db, provider_zone="Europe/Warsaw"):
    """Called only after ITN signature, service, order, currency and amount validation."""
    remote_id = transaction.findtext("remoteID", "")
    payment_date = transaction.findtext("paymentDate", "")
    if not re.fullmatch(r"[A-Za-z0-9]{1,20}", remote_id) or not re.fullmatch(r"[0-9]{14}", payment_date):
        raise HTTPException(400, "Niepoprawny identyfikator lub data wpłaty Autopay.")
    try:
        paid_at = parse_payment_date(payment_date, provider_zone)
        amount = Decimal(transaction.findtext("amount", ""))
        if not amount.is_finite() or amount <= 0 or amount != amount.quantize(Decimal("0.01")):
            raise InvalidOperation
        grosz = int(amount * 100)
        if grosz > 99999999:
            raise InvalidOperation
    except (ValueError, InvalidOperation):
        raise HTTPException(400, "Niepoprawna data lub kwota wpłaty Autopay.")
    acquire_payment_limit_lock(db)
    existing = db.query(AutopayReceipt).filter_by(
        environment=order.environment, service_id=order.service_id, remote_id=remote_id,
    ).first()
    if existing:
        if (existing.order_reference != order.id or existing.amount_grosz != grosz
                or existing.currency != order.currency or existing.provider_payment_date != payment_date):
            raise HTTPException(409, "Transakcja Autopay jest już zapisana z innymi danymi.")
        return existing
    payer = db.get(User, order.user_id) if order.user_id is not None else None
    entry = AutopayReceipt(
        order_id=order.id, order_reference=order.id, service_id=order.service_id,
        remote_id=remote_id, environment=order.environment, amount_grosz=grosz, currency=order.currency,
        estimated_fee_grosz=estimated_fee(grosz), fee_tariff=FEE_TARIFF,
        paid_on=paid_at.date(), paid_at=paid_at.isoformat(), provider_payment_date=payment_date,
        recorded_at=datetime.now(timezone.utc).isoformat(), package=order.package, period=order.period,
        payer_name=" ".join(filter(None, [payer.first_name, payer.last_name]))[:255] if payer else "",
        payer_email=(payer.email or "")[:255] if payer else "",
    )
    db.add(entry)
    db.flush()
    return entry


def receipt_query(db, environment):
    return db.query(AutopayReceipt).filter(AutopayReceipt.environment == environment, AutopayReceipt.currency == "PLN")


def year_query(query, year):
    return query.filter(AutopayReceipt.paid_on >= date(year, 1, 1), AutopayReceipt.paid_on < date(year + 1, 1, 1))


def totals(query):
    count, amount, fee = query.with_entities(func.count(AutopayReceipt.id),
        func.coalesce(func.sum(AutopayReceipt.amount_grosz), 0),
        func.coalesce(func.sum(AutopayReceipt.estimated_fee_grosz), 0)).one()
    return financial_totals(count, amount, fee)


def monthly_totals(query):
    month = extract("month", AutopayReceipt.paid_on)
    groups = query.with_entities(month, func.count(AutopayReceipt.id), func.sum(AutopayReceipt.amount_grosz),
                                 func.sum(AutopayReceipt.estimated_fee_grosz)).group_by(month).all()
    grouped = {int(number): (count, amount, fee) for number, count, amount, fee in groups}
    return [{"month": number, "label": label, **financial_totals(*grouped.get(number, (0, 0, 0)))}
            for number, label in enumerate(MONTHS, 1)]


def monthly_receipts(entries):
    # A downloaded report uses one immutable query result for every worksheet.
    # New ITNs arriving during export cannot change only its monthly totals.
    rows = [{"month": number, "label": label, "count": 0, "grosz": 0, "fee": 0} for number, label in enumerate(MONTHS, 1)]
    for entry in entries:
        row = rows[entry.paid_on.month - 1]
        row["count"] += 1
        row["grosz"] += entry.amount_grosz
        row["fee"] += entry.estimated_fee_grosz
    return [{"month": row["month"], "label": row["label"], **financial_totals(row["count"], row["grosz"], row["fee"])} for row in rows]


def quarterly_totals(months, year):
    limit = QUARTER_LIMITS_GROSZ.get(year)
    rows = []
    for quarter in range(1, 5):
        selected = months[(quarter - 1) * 3:quarter * 3]
        amount = sum(int(Decimal(row["amount"]) * 100) for row in selected)
        fee = sum(int(Decimal(row["estimated_fee"]) * 100) for row in selected)
        status = ("unknown" if limit is None else "exceeded" if amount > limit else
                  "reached" if amount == limit else "warning" if amount * 100 >= limit * 80 else "below")
        rows.append({"quarter": quarter, **financial_totals(sum(row["count"] for row in selected), amount, fee),
                     "limit": money(limit) if limit is not None else None,
                     "remaining": money(limit - amount) if limit is not None else None,
                     "used_percent": float(Decimal(amount * 100) / limit) if limit is not None else None,
                     "status": status})
    return rows


def serialize_receipt(entry):
    return {
        "id": entry.id, "paid_on": entry.paid_on.isoformat(), "paid_at": entry.paid_at,
        "recorded_at": entry.recorded_at, "payer_name": entry.payer_name, "payer_email": entry.payer_email,
        "description": receipt_description(entry.package, entry.period), "amount": money(entry.amount_grosz),
        "estimated_fee": money(entry.estimated_fee_grosz),
        "estimated_net": money(entry.amount_grosz - entry.estimated_fee_grosz), "fee_tariff": entry.fee_tariff,
        "currency": entry.currency, "remote_id": entry.remote_id, "service_id": entry.service_id,
        "order_reference": entry.order_reference, "environment": entry.environment,
    }


def payment_overview(db, year, environment, month=None, search="", page=1, page_size=50):
    all_receipts = receipt_query(db, environment)
    annual = year_query(all_receipts, year)
    matching = annual
    if month is not None:
        matching = matching.filter(extract("month", AutopayReceipt.paid_on) == month)
    if search.strip():
        escaped = search.strip().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        matching = matching.filter(or_(*[
            field.ilike(f"%{escaped}%", escape="\\") for field in (
                AutopayReceipt.payer_name, AutopayReceipt.payer_email,
                AutopayReceipt.remote_id, AutopayReceipt.order_reference,
            )
        ]))
    matching_totals = totals(matching)
    rows = matching.order_by(AutopayReceipt.paid_on.desc(), AutopayReceipt.paid_at.desc(), AutopayReceipt.id.desc()).offset((page - 1) * page_size).limit(page_size).all()
    years = [int(row[0]) for row in all_receipts.with_entities(extract("year", AutopayReceipt.paid_on)).distinct().all()]
    months = monthly_totals(annual)
    return {
        "year": year, "environment": environment, "page": page, "page_size": page_size,
        "years": sorted(set(years + [datetime.now(WARSAW).year]), reverse=True),
        "annual": totals(annual), "lifetime": totals(all_receipts), "matching": matching_totals,
        "months": months, "quarters": quarterly_totals(months, year), "entries": [serialize_receipt(row) for row in rows],
        "report_note": REPORT_NOTE, "tax_source_url": TAX_SOURCE,
        "fee_note": FEE_NOTE, "limit_note": LIMIT_NOTE, "limit_source_url": LIMIT_SOURCE,
        "payment_guard": payment_limit_state(db),
    }


def export_row(entry):
    return [entry.paid_at, entry.payer_name, entry.payer_email, receipt_description(entry.package, entry.period),
            money(entry.amount_grosz), entry.remote_id, entry.service_id, entry.order_reference, entry.recorded_at,
            money(entry.estimated_fee_grosz), money(entry.amount_grosz - entry.estimated_fee_grosz), entry.fee_tariff]


def csv_text(value):
    text = str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text


def build_payments_csv(entries, year):
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(HEADERS)
    for entry in entries:
        row = export_row(entry)
        row[4] = row[4].replace(".", ",")
        row[9] = row[9].replace(".", ",")
        row[10] = row[10].replace(".", ",")
        writer.writerow([value if index in (4, 9, 10) else csv_text(value) for index, value in enumerate(row)])
    amount = sum(e.amount_grosz for e in entries)
    fee = sum(e.estimated_fee_grosz for e in entries)
    writer.writerow([f"SUMA WPŁAT {year}", "", "", "", money(amount).replace(".", ","), "", "", "", "",
                     money(fee).replace(".", ","), money(amount - fee).replace(".", ","), ""])
    writer.writerow([REPORT_NOTE])
    writer.writerow([FEE_NOTE])
    writer.writerow([LIMIT_NOTE])
    return output.getvalue().encode("utf-8-sig")


def write_excel_row(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
            cell.data_type = "s"


def build_payments_workbook(entries, year, months, seller):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Podsumowanie"
    for row in [
        ["Roczne zestawienie rzeczywistych wpłat Autopay — pomocniczo do PIT-36"],
        ["Rok", year], ["Sprzedawca", seller.get("name", "")], ["Adres", seller.get("address", "")],
        ["Wygenerowano (Polska)", datetime.now(WARSAW).isoformat()],
        ["Liczba potwierdzonych wpłat", len(entries)],
        ["Suma wpłat klientów [PLN], przed prowizjami i zwrotami", float(Decimal(sum(e.amount_grosz for e in entries)) / 100)],
        ["Prowizje Autopay — szacunek [PLN]", float(Decimal(sum(e.estimated_fee_grosz for e in entries)) / 100)],
        ["Zwroty klientom", "Nie uwzględniono"],
        ["Zakres raportu", REPORT_NOTE], ["Zasady PIT-36 — Ministerstwo Finansów", TAX_SOURCE],
        ["Po prowizjach — szacunek [PLN]", float(Decimal(sum(e.amount_grosz - e.estimated_fee_grosz for e in entries)) / 100)],
        ["Zasady szacunku prowizji", FEE_NOTE], ["Kontrola limitu — zakres", LIMIT_NOTE],
        ["Zasady limitu — Ministerstwo Finansów", LIMIT_SOURCE], ["Inne koszty", "Nie uwzględniono"],
    ]:
        write_excel_row(summary, row)
    summary["B7"].number_format = '#,##0.00 "zł"'
    summary["B8"].number_format = '#,##0.00 "zł"'
    summary["B12"].number_format = '#,##0.00 "zł"'
    summary.column_dimensions["A"].width = 65
    summary.column_dimensions["B"].width = 100
    for number in (10, 13, 14):
        summary.cell(number, 2).alignment = Alignment(wrap_text=True)
        summary.row_dimensions[number].height = 80
    monthly = workbook.create_sheet("Miesiące")
    write_excel_row(monthly, ["Miesiąc", "Liczba wpłat", "Suma wpłat [PLN]", "Prowizje — szacunek [PLN]", "Po prowizjach — szacunek [PLN]"])
    for month in months:
        write_excel_row(monthly, [month["label"], month["count"], *[float(Decimal(month[key])) for key in ("amount", "estimated_fee", "estimated_net")]])
        for column in (3, 4, 5):
            monthly.cell(monthly.max_row, column).number_format = '#,##0.00 "zł"'
    monthly.column_dimensions["A"].width = 22
    monthly.column_dimensions["B"].width = 20
    monthly.column_dimensions["C"].width = 25
    monthly.column_dimensions["D"].width = 30
    monthly.column_dimensions["E"].width = 35
    detail = workbook.create_sheet("Wpłaty")
    write_excel_row(detail, HEADERS)
    for entry in entries:
        row = export_row(entry)
        row[4] = float(Decimal(row[4]))
        row[9] = float(Decimal(row[9]))
        row[10] = float(Decimal(row[10]))
        write_excel_row(detail, row)
        for column in (5, 10, 11):
            detail.cell(detail.max_row, column).number_format = '#,##0.00 "zł"'
    for column, width in zip("ABCDEFGHIJKL", [30, 28, 35, 45, 22, 25, 15, 36, 35, 30, 35, 35]):
        detail.column_dimensions[column].width = width
    quarterly = workbook.create_sheet("Kwartały")
    write_excel_row(quarterly, ["Kwartał", "Liczba wpłat", "Wpłaty przed prowizją [PLN]", "Prowizje — szacunek [PLN]",
                              "Po prowizjach — szacunek [PLN]", "Limit przychodu należnego [PLN]", "Margines wg wpłat [PLN]"])
    for item in quarterly_totals(months, year):
        write_excel_row(quarterly, [f"Q{item['quarter']} {year}", item["count"],
                        *[float(Decimal(item[key])) for key in ("amount", "estimated_fee", "estimated_net")],
                        *[float(Decimal(item[key])) if item[key] is not None else "Nie ustalono limitu dla tego roku"
                          for key in ("limit", "remaining")]])
        for column in range(3, 8):
            quarterly.cell(quarterly.max_row, column).number_format = '#,##0.00 "zł"'
    for column in "ABCDEFG":
        quarterly.column_dimensions[column].width = 35
    write_excel_row(quarterly, [LIMIT_NOTE])
    quarterly.auto_filter.ref = "A1:G5"
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        if sheet not in (summary, quarterly):
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="14532D")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_payments_router(get_db, get_current_admin, seller_details):
    router = APIRouter(prefix="/admin/payments", tags=["Wpłaty Autopay"])

    @router.get("")
    def overview(
        response: Response, year: int | None = Query(default=None, ge=2000, le=2100),
        environment: Literal["production"] = "production",
        month: int | None = Query(default=None, ge=1, le=12),
        search: str = Query(default="", max_length=100), page: int = Query(default=1, ge=1),
        page_size: int = Query(default=50, ge=1, le=100),
        admin=Depends(get_current_admin), db=Depends(get_db),
    ):
        response.headers["Cache-Control"] = "no-store"
        return payment_overview(db, year or datetime.now(WARSAW).year, environment, month, search, page, page_size)

    @router.get("/report")
    def annual_report(
        year: int = Query(ge=2000, le=2100), format: Literal["xlsx", "csv"] = "xlsx",
        environment: Literal["production"] = "production",
        admin=Depends(get_current_admin), db=Depends(get_db),
    ):
        query = year_query(receipt_query(db, environment), year)
        entries = query.order_by(AutopayReceipt.paid_on, AutopayReceipt.paid_at, AutopayReceipt.id).all()
        if format == "csv":
            content = build_payments_csv(entries, year)
            media_type = "text/csv; charset=utf-8"
        else:
            content = build_payments_workbook(entries, year, monthly_receipts(entries), seller_details())
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return Response(content, media_type=media_type, headers={
            "Content-Disposition": f'attachment; filename="autopay-wplaty-{year}.{format}"',
            "Cache-Control": "no-store",
        })

    return router
